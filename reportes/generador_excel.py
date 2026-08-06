from openpyxl import load_workbook
from database.operaciones import gastos_mes

def generar_reporte(mes,ano):
    try: 
        # PASO 1 DICCIONARIOS DE LOS MESES
        meses = {
            "01": "ENERO", "02": "FEBRERO", "03": "MARZO",
            "04": "ABRIL", "05": "MAYO", "06": "JUNIO",
            "07": "JULIO", "08": "AGOSTO", "09": "SEPTIEMBRE",
            "10": "OCTUBRE", "11": "NOVIEMBRE", "12": "DICIEMBRE"
        }
        # PASO 2 ABRO EL LIBRO INDICADO 
        wb = load_workbook("Reembolsos.xlsx")
        
        # PASO 3 CONDICIONAL SI EXISTE LA HOJA
        nombre_mes = meses[mes]
        nombre_hoja = f"{nombre_mes} {ano}"

        if nombre_hoja in wb.sheetnames:
            hoja = wb[nombre_hoja]
        else:
            hoja = wb.copy_worksheet(wb.active)
            hoja.title = nombre_hoja
        # PASO 4 Y 5 CONSULTAR LA BASE DE DATOS CON GASTOS_MES
        filas = gastos_mes(mes, ano)
        fila_contable = 6
        fila_no_contable=25
        for fila in filas:
            tipo = fila[8]
            if tipo == "contable":
                hoja.cell(row=fila_contable, column=1).value = fila[1]
                hoja.cell(row=fila_contable, column=2).value = fila[2]
                hoja.cell(row=fila_contable, column=3).value = fila[3]
                hoja.cell(row=fila_contable, column=4).value = fila[4]
                hoja.cell(row=fila_contable, column=5).value = fila[5]
                hoja.cell(row=fila_contable, column=6).value = fila[6]
                hoja.cell(row=fila_contable, column=7).value = fila[7]
                fila_contable += 1
            elif tipo == "no contable":
                hoja.cell(row=fila_no_contable, column=1).value = fila[1]
                hoja.cell(row=fila_no_contable, column=2).value = fila[2]
                hoja.cell(row=fila_no_contable, column=3).value = fila[3]
                hoja.cell(row=fila_no_contable, column=4).value = fila[4]
                hoja.cell(row=fila_no_contable, column=5).value = fila[5]
                hoja.cell(row=fila_no_contable, column=6).value = fila[6]
                hoja.cell(row=fila_no_contable, column=7).value = fila[7]
                fila_no_contable += 1
        wb.save("Reembolsos.xlsx")
        return "Reembolsos.xlsx"
    except Exception as e:
        print(f"Error escribiendo el excel: {e}")
        return False